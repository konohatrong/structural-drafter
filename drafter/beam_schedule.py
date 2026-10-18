"""
The office RC beam schedule (RC rules §15.3, row form; user rules 2026-10-18), driven by a spreadsheet.

One row per mark: NO. | SIZE (mm) | SECTIONS in five fixed columns END SUPPORT | MID SPAN | CONTINUOUS SUPPORT |
CANTILEVER | ALL SPAN (an unused one left empty); under each section a TOP / BOT / STIRRUP / SIDE table; the section
drawn with the slab as a dummy flange and EIT break lines. Drawn in model space at full size (mm), shown at 1:25.

Spreadsheet (.xlsx, sheet BEAMS), one row per mark and section (user, 2026-10-18; rules:
docs/concrete/BEAM_SCHEDULE_RULES.md):
  MARK | b | h | SLAB t | SECTION | TOP | TOP DB | BOT | BOT DB | STIRRUP | SIDE | NOTE
  TOP / BOT, layers from that face, outer layer first:  '/' separates layers;  '+' starts the extra bars of the section,
  each '/' after it a further inner layer;  N(+M) is one layer of N nominal and M extra bars:
    3/3/3   9 nominal in 3 layers                         drawn "9 - DB16"
    3+2     3 nominal (layer 1) + 2 extra (layer 2)        drawn "3+2 - DB16"
    3/3+2   6 nominal (layers 1, 2) + 2 extra (layer 3)    drawn "6+2 - DB16"
    3+3/2   3 nominal + 5 extra (layers 2, 3)              drawn "3+5 - DB16"
    2(+1)   one layer: 2 nominal + 1 extra                 drawn "2+1 - DB16"
  STIRRUP "2-RB9@100": legs - bar - spacing; 2 legs = one hoop, 3 = hoop + one crosstie, 4 = hoop + an inner hoop;
  SLAB t: the slab thickness drawn as the dummy flange (0 or blank: no slab); a blank MARK / b / h / SLAB t repeats the
  row above. Sheet SETTINGS: cover, clear distance between layers, flange length drawn, hook length.
On the drawing the bars read nominal + extra ("6+2 - DB16") and the stirrups as entered.

API:  read_xlsx(path) -> (marks, settings);  write_xlsx(path, marks, settings=None);  draw(sp, marks, settings, S=25)
CLI:  python -m drafter.beam_schedule template <folder>     the office template (.xlsx, .dxf and, with AutoCAD, .dwg)
      python -m drafter.beam_schedule draw <in.xlsx> <out.dxf>
"""
import math
import re
from pathlib import Path

TAGS = ["END SUPPORT", "MID SPAN", "CONTINUOUS SUPPORT", "CANTILEVER", "ALL SPAN"]
LABS = ("TOP", "BOT", "STIRRUP", "SIDE")
HEAD = ["MARK", "b (mm)", "h (mm)", "SLAB t (mm)", "SECTION", "TOP", "TOP DB", "BOT", "BOT DB", "STIRRUP", "SIDE", "NOTE"]
INSTRUCTIONS = [
    ["RC BEAM SCHEDULE - HOW TO FILL THIS FILE (office standard, docs/concrete/BEAM_SCHEDULE_RULES.md)"],
    [],
    ["SHEET BEAMS"],
    ["one row", "", "per beam mark and section. Several sections of one mark: leave MARK, b, h, SLAB t blank on the "
                    "rows below the first - the row above repeats"],
    ["MARK", "", "beam mark as on the plans: B1, B1C (cantilever), RB1 (roof), B5A (all-span bars)"],
    ["b, h", "", "width and depth, mm"],
    ["SLAB t", "", "slab thickness drawn as the dummy flange, mm; 0 = no slab at that level"],
    ["SECTION", "", "END SUPPORT, MID SPAN, CONTINUOUS SUPPORT, CANTILEVER or ALL SPAN (pick from the list)"],
    ["TOP", "", "top bars, layers from the top fibre (the rules below)"],
    ["BOT", "", "bottom bars, layers from the bottom fibre (the rules below)"],
    ["TOP DB, BOT DB", "", "bar size of that face, mm (16 = DB16)"],
    ["STIRRUP", "", "legs-bar@spacing (the rules below)"],
    ["SIDE", "", "side-face bars, each face: 2-DB12; - = none"],
    ["NOTE", "", "free text (not drawn)"],
    [],
    ["BAR INPUT RULES (TOP and BOT)"],
    ["/", "", "separates layers, outer layer first (at the face) - the last one is the innermost layer"],
    ["+", "", "starts the EXTRA bars of this section (additional to the nominal, continuous bars); each / after it is "
              "a further inner layer of extra bars"],
    ["N(+M)", "", "one layer with N nominal and M extra bars sharing it"],
    ["example", "input", "meaning / drawn as"],
    ["", "3/3/3", "9 nominal bars in 3 layers, no extra: 9 - DB16"],
    ["", "3+2", "3 nominal (layer 1) + 2 extra (layer 2): 3+2 - DB16"],
    ["", "3/3+2", "6 nominal (layers 1, 2) + 2 extra (layer 3): 6+2 - DB16"],
    ["", "3+3/2", "3 nominal (layer 1) + 5 extra (layers 2, 3): 3+5 - DB16"],
    ["", "2(+1)", "one layer: 2 nominal + 1 extra: 2+1 - DB16"],
    ["", "3/2(+1)+2", "5 nominal + 3 extra: layer 1 3 nominal, layer 2 2 nominal + 1 extra, layer 3 2 extra: 5+3 - DB16"],
    ["", "2", "2 nominal bars in one layer: 2 - DB16"],
    ["write as text", "", "type the bar cells as text (the cells are formatted @) so Excel does not turn 3/3 into a date"],
    [],
    ["STIRRUP INPUT RULES"],
    ["n-RBd@s", "", "n legs, round bar RB d mm (or DB), spacing s mm"],
    ["", "2-RB9@100", "2 legs = one closed hoop"],
    ["", "3-RB9@100", "3 legs = hoop + one crosstie (on a bar present in both the top and bottom outer layers)"],
    ["", "4-RB9@100", "4 legs = hoop + an inner hoop (on two such bars)"],
    ["", "RB9@200", "no leg count = 2 legs"],
    [],
    ["SHEET SETTINGS"],
    ["cover", "", "mm to the stirrup; layer_clear: mm between layers; flange: mm of slab drawn each side; hook: mm "
                  "(0 = 6 db + 25)"],
    [],
    ["AFTER EDITING"],
    ["", "", "save the file and tell the drafter (or run the project's build): the schedule sheet and the strength "
             "check are both made from this file"],
]
SETTINGS = {"cover": (40, "mm, to the stirrup (office general notes TABLE 7: beams 40)"),
            "layer_clear": (25, "mm clear between bar layers (ACI 25.2.2: 25 mm or db)"),
            "flange": (300, "mm of slab dummy drawn each side of the beam"),
            "hook": (0, "mm, stirrup hook extension drawn; 0 = 6 db + 25")}
AREA = {6: 28.3, 9: 63.6, 10: 78.5, 12: 113.1, 16: 201.1, 20: 314.2, 25: 490.9, 28: 615.8, 32: 804.2}
_TOK = re.compile(r"^(\d+)(?:\(\+(\d+)\))?$")                             # N or N(+M)
_ST = re.compile(r"^\s*(?:(\d+)\s*-\s*)?(RB|DB)(\d+)\s*@\s*(\d+)\s*$", re.I)


# ------------------------------------------------------------------------------------------- the notation
def _split_plus(s):
    """split at the first '+' outside parentheses: '3/2(+1)+2' -> ('3/2(+1)', '2')"""
    depth = 0
    for k, ch in enumerate(s):
        depth += ch == "("
        depth -= ch == ")"
        if ch == "+" and depth == 0:
            return s[:k], s[k + 1:]
    return s, None


def parse_bars(s):
    """the bar grammar -> layers, outer first: [{'n': nominal, 'x': extra}, ...]; '' / '-' / None -> []
    '3/3+2' -> [{n 3}, {n 3}, {x 2}];  '2(+1)' -> [{n 2, x 1}];  '3+3/2' -> [{n 3}, {x 3}, {x 2}]"""
    if s is None or str(s).strip() in ("", "-"):
        return []
    s = str(s).replace(" ", "")
    if s.replace(".", "", 1).isdigit() and "." in s:                 # Excel turned '3' into '3.0'
        s = str(int(float(s)))
    nom, ext = _split_plus(s)
    bad = ValueError(f"bars '{s}': layers with '/', extra bars after '+', a shared layer as N(+M) - e.g. 3/3+2, 2(+1)")
    out = []
    for tok in nom.split("/"):
        m = _TOK.match(tok)
        if not m:
            raise bad
        out.append({"n": int(m.group(1)), "x": int(m.group(2) or 0)})
    if ext is not None:
        for tok in ext.split("/"):
            if not tok.isdigit():
                raise bad
            out.append({"n": 0, "x": int(tok)})
    if any(l["n"] < 1 for l in out[:len(nom.split("/"))]) or any(l["x"] < 1 for l in out[len(nom.split("/")):]):
        raise ValueError(f"bars '{s}': every layer needs at least one bar")
    return out


def fmt_bars(layers):
    """layers -> the grammar ('3/3+2', '2(+1)')"""
    nom = [f"{l['n']}" + (f"(+{l['x']})" if l["x"] else "") for l in layers if l["n"]]
    ext = [f"{l['x']}" for l in layers if not l["n"]]
    return "/".join(nom) + ("+" + "/".join(ext) if ext else "")


def totals(layers):
    """bars per layer (nominal + extra), outer first"""
    return [l["n"] + l["x"] for l in layers]


def parse_stirrup(s):
    """'2-RB9@100' -> {'legs': 2, 'mark': 'RB', 'db': 9, 's': 100}; legs default 2"""
    m = _ST.match(str(s or ""))
    if not m:
        raise ValueError(f"stirrup '{s}': write legs-bar@spacing, e.g. 2-RB9@100")
    legs = int(m.group(1) or 2)
    if legs < 2:
        raise ValueError(f"stirrup '{s}': at least 2 legs (one hoop)")
    return {"legs": legs, "mark": m.group(2).upper(), "db": int(m.group(3)), "s": int(m.group(4))}


def fmt_stirrup(st):
    return f"{st['legs']}-{st['mark']}{st['db']}@{st['s']}"


def bars_label(layers, db):
    """the drawing text: nominal + extra, '6+2 - DB16'; no extra: '9 - DB16'"""
    n, x = sum(l["n"] for l in layers), sum(l["x"] for l in layers)
    return f"{n}+{x} - DB{db}" if x else f"{n} - DB{db}"


def bar_area(layers, db):
    return sum(totals(layers)) * AREA[db]


# ------------------------------------------------------------------------------------------- the spreadsheet
def _int(v, what, row):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        raise ValueError(f"row {row}: {what} '{v}' is not a number")


def read_xlsx(path):
    """the BEAMS and SETTINGS sheets -> (marks, settings); marks = {mark: {b, h, slab, cols: {tag: {...}}}}"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["BEAMS"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and str(r[0] or "").strip().upper() == "MARK")
    head = [str(v or "").strip().upper() for v in rows[hi]]
    col = {h: head.index(h.upper()) for h in HEAD}
    marks, last = {}, {}
    for i, r in enumerate(rows[hi + 1:], start=hi + 2):
        get = lambda h: r[col[h]] if col[h] < len(r) else None       # noqa: E731
        if not any(v not in (None, "") for v in r):
            continue
        cur = {k: get(k) if get(k) not in (None, "") else last.get(k) for k in ("MARK", "b (mm)", "h (mm)", "SLAB t (mm)")}
        last = cur
        mk = str(cur["MARK"]).strip()
        tag = str(get("SECTION") or "").strip().upper()
        if tag not in TAGS:
            raise ValueError(f"row {i}: SECTION '{tag}' is not one of {TAGS}")
        m = marks.setdefault(mk, {"b": _int(cur["b (mm)"], "b", i), "h": _int(cur["h (mm)"], "h", i),
                                  "slab": _int(cur["SLAB t (mm)"] or 0, "SLAB t", i), "cols": {}})
        if tag in m["cols"]:
            raise ValueError(f"row {i}: {mk} {tag} given twice")
        try:
            m["cols"][tag] = {"top": {"layers": parse_bars(get("TOP")), "db": _int(get("TOP DB"), "TOP DB", i)},
                              "bot": {"layers": parse_bars(get("BOT")), "db": _int(get("BOT DB"), "BOT DB", i)},
                              "stirrup": parse_stirrup(get("STIRRUP")),
                              "side": str(get("SIDE") or "-").strip() or "-", "note": str(get("NOTE") or "").strip()}
        except ValueError as e:
            raise ValueError(f"row {i} ({mk} {tag}): {e}")
    settings = {k: v[0] for k, v in SETTINGS.items()}
    if "SETTINGS" in wb.sheetnames:
        for r in wb["SETTINGS"].iter_rows(min_row=2, values_only=True):
            if r and r[0] in settings and r[1] not in (None, ""):
                settings[r[0]] = float(r[1])
    return marks, settings


def write_xlsx(path, marks, settings=None, title="RC BEAM SCHEDULE"):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BEAMS"
    ws.append([title + " - one row per mark and section; how to fill: sheet INSTRUCTIONS"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(HEAD)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
        c.alignment = Alignment(horizontal="center")
    for mk, m in marks.items():
        first = True
        for tag in TAGS:
            c = m["cols"].get(tag)
            if not c:
                continue
            ws.append([mk if first else None, m["b"] if first else None, m["h"] if first else None,
                       (m.get("slab") or 0) if first else None, tag, fmt_bars(c["top"]["layers"]), c["top"]["db"],
                       fmt_bars(c["bot"]["layers"]), c["bot"]["db"], fmt_stirrup(c["stirrup"]), c.get("side", "-"),
                       c.get("note", "")])
            first = False
    for k, w in zip("ABCDEFGHIJKL", (8, 8, 8, 10, 22, 10, 8, 10, 8, 14, 10, 40)):
        ws.column_dimensions[k].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row[5:10]:
            c.number_format = "@"                                     # 3/3/2 stays text, not a date
    dv = DataValidation(type="list", formula1='"' + ",".join(TAGS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("E3:E500")
    ws.freeze_panes = "A3"
    nt = wb.create_sheet("INSTRUCTIONS")
    for line in INSTRUCTIONS:
        nt.append(line)
    nt["A1"].font = Font(bold=True, size=12)
    for r in nt.iter_rows():
        if r[0].value and len(r) > 1 and r[1].value is None and r[0].row > 1:
            r[0].font = Font(bold=True)
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 26
    nt.column_dimensions["B"].width = 24
    nt.column_dimensions["C"].width = 90
    st = wb.create_sheet("SETTINGS")
    st.append(["NAME", "VALUE", "NOTE"])
    for k, (v, note) in SETTINGS.items():
        st.append([k, (settings or {}).get(k, v), note])
    st.column_dimensions["A"].width = 14
    st.column_dimensions["C"].width = 70
    wb.save(path)


# ------------------------------------------------------------------------------------------- the drawing
W_NO, W_SIZE, W_SEC, W_LAB = 14.0, 20.0, 46.0, 12.0                    # paper mm
H_HEAD, H_SUB, TH = 6.0, 5.0, 2.0


def _bar_x(n, x0, x1):
    return [(x0 + x1) / 2] if n == 1 else [x0 + (x1 - x0) * k / (n - 1) for k in range(n)]


def _inner_legs(top_xs, bot_xs, n_inner):
    """x of inner legs: bar positions in both outer layers, nearest the middle (RC rules §15.3: never a standalone leg)"""
    xc = (top_xs[0] + top_xs[-1]) / 2
    common = [x for x in top_xs[1:-1] if any(abs(x - y) < 15 for y in bot_xs[1:-1])]
    common.sort(key=lambda x: abs(x - xc))
    return sorted(common[:n_inner]) if len(common) >= n_inner else None


def draw_section(sp, cx, ytop, b, h, c, set_, slab, S, warn):
    """one section, beam top centre at (cx, ytop), model mm"""
    from drafter.td_engine import pline, dot, rdot, stirrup, crosstie, zbreak
    cov, clear, fl = set_["cover"], set_["layer_clear"], set_["flange"]
    st = c["stirrup"]
    ds = st["db"]
    hook = set_["hook"] or 6 * ds + 25
    x0, y0, yt_ = cx - b / 2, ytop - h, ytop
    if slab:
        xl, xr = x0 - fl, x0 + b + fl
        pline(sp, [(xl, yt_), (xr, yt_)], "S-CONC")
        pline(sp, [(xr, yt_ - slab), (x0 + b, yt_ - slab), (x0 + b, y0), (x0, y0), (x0, yt_ - slab), (xl, yt_ - slab)],
              "S-CONC")
        for x in (xl, xr):
            zbreak(sp, (x, yt_ + 2.0 * S), (x, yt_ - slab - 2.0 * S), S)
    else:
        pline(sp, [(x0, y0), (x0 + b, y0), (x0 + b, yt_), (x0, yt_)], "S-CONC", close=True)
    T = {"db": c["top"]["db"], "layers": totals(c["top"]["layers"])}
    B = {"db": c["bot"]["db"], "layers": totals(c["bot"]["layers"])}
    rt, rb = rdot(T["db"], S), rdot(B["db"], S)
    xt0, xt1 = x0 + cov + ds + T["db"] / 2, x0 + b - cov - ds - T["db"] / 2
    xb0, xb1 = x0 + cov + ds + B["db"] / 2, x0 + b - cov - ds - B["db"] / 2
    yt = ytop - cov - ds - T["db"] / 2
    yb = y0 + cov + ds + B["db"] / 2
    top_xs = _bar_x(T["layers"][0], xt0, xt1) if T["layers"] else [xt0, xt1]
    bot_xs = _bar_x(B["layers"][0], xb0, xb1) if B["layers"] else [xb0, xb1]
    for k, n in enumerate(T["layers"]):
        for x in (top_xs if k == 0 else _bar_x(n, xt0, xt1)):
            dot(sp, (x, yt - k * (T["db"] + clear)), rt)
    for k, n in enumerate(B["layers"]):
        for x in (bot_xs if k == 0 else _bar_x(n, xb0, xb1)):
            dot(sp, (x, yb + k * (B["db"] + clear)), rb)
    Rt, Rb = rt + ds / 2, rb + ds / 2
    stirrup(sp, (xt0, yt, Rt), (xt1, yt, Rt), (xb1, yb, Rb), (xb0, yb, Rb), hook)
    sm = re.match(r"^\s*(\d+)\s*-\s*DB(\d+)", str(c.get("side") or ""), re.I)
    if sm:                                                            # side-face bars, n per face, inside the legs
        n, dbs = int(sm.group(1)), int(sm.group(2))
        y_hi = yt - (len(T["layers"]) - 1) * (T["db"] + clear)
        y_lo = yb + (len(B["layers"]) - 1) * (B["db"] + clear)
        for k in range(1, n + 1):
            y = y_hi + (y_lo - y_hi) * k / (n + 1)
            for x in (x0 + cov + ds + dbs / 2, x0 + b - cov - ds - dbs / 2):
                dot(sp, (x, y), rdot(dbs, S))
    if st["legs"] > 2:
        n_in = st["legs"] - 2
        if n_in == 1:
            xs = _inner_legs(top_xs, bot_xs, 1)
            if xs is None:
                warn(f"beam schedule: {st['legs']} legs but no bar in both outer layers for the crosstie "
                     f"(top {len(top_xs)}, bottom {len(bot_xs)} bars): rearrange the bars")
            else:
                crosstie(sp, (xs[0], yt), (xs[0], yb), max(Rt, Rb), hook)
        elif n_in == 2:
            xs = _inner_legs(top_xs, bot_xs, 2)
            if xs is None:
                warn(f"beam schedule: 4 legs but no two bars in both outer layers for the inner hoop: rearrange")
            else:
                stirrup(sp, (xs[0], yt, Rt), (xs[1], yt, Rt), (xs[1], yb, Rb), (xs[0], yb, Rb), hook)
        else:
            warn(f"beam schedule: {st['legs']} legs not drawn (up to 4)")


def draw(sp, marks, settings=None, S=25, txt_layer="S-TXT_2.5_MM", lines=("S-TTLB", "S-TTLB-THIN"), warn=None):
    """the schedule at full size in model space, its top-left corner at (0, 0); returns the marks drawn"""
    from drafter.td_engine import text, line, TA
    from drafter import td_engine
    warn = warn or td_engine.warn
    set_ = {k: v[0] for k, v in SETTINGS.items()}
    set_.update(settings or {})
    LT, LB = lines
    t = lambda s, x, y, al=TA.MIDDLE_CENTER, st="AN": text(sp, s, (x, y), TH * S, txt_layer, al, style=st)  # noqa
    x_size, x_sec = W_NO * S, (W_NO + W_SIZE) * S
    xs = [x_sec + k * W_SEC * S for k in range(len(TAGS) + 1)]
    W, hh, y = xs[-1], H_HEAD * S, 0.0
    line(sp, (0, y), (W, y), LT)
    t("NO.", W_NO * S / 2, y - hh, st="ANB")
    t("SIZE (mm)", x_size + W_SIZE * S / 2, y - hh, st="ANB")
    t("SECTIONS", (x_sec + W) / 2, y - hh / 2, st="ANB")
    line(sp, (x_sec, y - hh), (W, y - hh), LB)
    for k, tag in enumerate(TAGS):
        t(tag, (xs[k] + xs[k + 1]) / 2, y - 1.5 * hh, st="ANB")
        if k:
            line(sp, (xs[k], y - hh), (xs[k], y - 2 * hh), LB)
    y -= 2 * hh
    line(sp, (0, y), (W, y), LT)
    done = []
    for mk, m in marks.items():
        b, h, slab = m["b"], m["h"], m.get("slab") or 0
        band = h + 2 * 4.0 * S
        rh = band + 4 * H_SUB * S
        top = y
        t(mk, W_NO * S / 2, top - rh / 2, st="ANB")
        t(f"{b} x {h}", x_size + W_SIZE * S / 2, top - rh / 2)
        ysub = top - band
        line(sp, (x_sec, ysub), (W, ysub), LB)
        for k, tag in enumerate(TAGS):
            xa, xb = xs[k], xs[k + 1]
            for j in range(1, 4):
                line(sp, (xa, ysub - j * H_SUB * S), (xb, ysub - j * H_SUB * S), LB)
            line(sp, (xa + W_LAB * S, ysub), (xa + W_LAB * S, ysub - 4 * H_SUB * S), LB)
            c = m["cols"].get(tag)
            vals = None if c is None else [bars_label(c["top"]["layers"], c["top"]["db"]),
                                           bars_label(c["bot"]["layers"], c["bot"]["db"]), fmt_stirrup(c["stirrup"]),
                                           c.get("side") or "-"]
            for j, lab in enumerate(LABS):
                yy = ysub - (j + 0.5) * H_SUB * S
                t(lab, xa + W_LAB * S / 2, yy)
                if vals:
                    t(vals[j], (xa + W_LAB * S + xb) / 2, yy)
            if c is not None:
                draw_section(sp, (xa + xb) / 2, top - 4.0 * S, b, h, c, set_, slab, S, warn)
        y -= rh
        line(sp, (0, y), (W, y), LB)
        done.append(mk)
    line(sp, (0, y), (W, y), LT)
    for x, lay in ((0, LT), (W, LT), (x_size, LB), (x_sec, LB)):
        line(sp, (x, 0), (x, y), lay)
    for x in xs[1:-1]:
        line(sp, (x, -2 * hh), (x, y), LB)
    return done


# ------------------------------------------------------------------------------------------- the office template
def _c(top, tdb, bot, bdb, st, side="-"):
    return {"top": {"layers": parse_bars(top), "db": tdb}, "bot": {"layers": parse_bars(bot), "db": bdb},
            "stirrup": parse_stirrup(st), "side": side}


EXAMPLE = {
    "B1": {"b": 250, "h": 600, "slab": 150, "cols": {
        "END SUPPORT": _c("3+2", 20, "2", 20, "2-RB9@100"),
        "MID SPAN": _c("3", 20, "2(+1)", 20, "2-RB9@200"),
        "CONTINUOUS SUPPORT": _c("3+2", 20, "2", 20, "2-RB9@100")}},
    "B2": {"b": 300, "h": 700, "slab": 150, "cols": {
        "END SUPPORT": _c("5+3", 20, "3", 20, "3-RB9@100", "2-DB12"),
        "MID SPAN": _c("3", 20, "3(+2)/2", 20, "3-RB9@200", "2-DB12"),
        "CONTINUOUS SUPPORT": _c("5/3+2", 20, "3", 20, "3-RB9@100", "2-DB12")}},
    "B1C": {"b": 250, "h": 600, "slab": 150, "cols": {
        "CANTILEVER": _c("3/2", 20, "2", 20, "2-RB9@100")}},
    "RB1": {"b": 250, "h": 400, "slab": 0, "cols": {
        "ALL SPAN": _c("2", 16, "2", 16, "2-RB9@150")}},
}


def template(folder):
    """the office template: BEAM_SCHEDULE_TEMPLATE.xlsx (example rows), .dxf (the schedule drawn from it at full size,
    for 1:25) and .dwg when AutoCAD's Core Console is installed"""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    xlsx = folder / "BEAM_SCHEDULE_TEMPLATE.xlsx"
    write_xlsx(xlsx, EXAMPLE, title="RC BEAM SCHEDULE - EXAMPLE ROWS (REPLACE)")
    marks, settings = read_xlsx(xlsx)
    dxf = folder / "BEAM_SCHEDULE_TEMPLATE.dxf"
    to_dxf(marks, settings, dxf, title="OFFICE RC BEAM SCHEDULE - TEMPLATE (RC RULES 15.3). MODEL SPACE AT FULL SIZE "
                                       "(mm), PLOT AT 1:25. DATA: BEAM_SCHEDULE_TEMPLATE.xlsx")
    dwg = to_dwg(dxf)
    return xlsx, dxf, dwg


def to_dxf(marks, settings, out, title=None, S=25, txt_layer="S-TEXT"):
    """the schedule alone in a new engine drawing (model space, full size) -> out"""
    from drafter import td_engine
    from drafter.td_engine import text, TA
    draw(td_engine.msp, marks, settings, S=S, txt_layer=txt_layer)
    if title:
        text(td_engine.msp, title, (0, 4.0 * S), 2.8 * S, txt_layer, TA.BOTTOM_LEFT, style="ANB")
    if td_engine.WARNINGS:
        raise SystemExit("!! " + "\n!! ".join(td_engine.WARNINGS))
    td_engine.doc.saveas(out)
    return out


def to_dwg(dxf):
    """SAVEAS 2018 DWG through the Core Console; None without AutoCAD"""
    import shutil
    import subprocess
    import tempfile
    from drafter import acad
    if not Path(acad.ACC).exists():
        return None
    tmp = Path(tempfile.mkdtemp())
    src, out = tmp / "s.dxf", tmp / "s.dwg"
    shutil.copy2(dxf, src)
    (tmp / "x.scr").write_bytes(f'FILEDIA\r\n0\r\n_.SAVEAS\r\n2018\r\n"{out}"\r\n'.encode("ascii"))
    subprocess.run([acad.ACC, "/i", str(src), "/s", str(tmp / "x.scr"), "/l", "en-US"], capture_output=True, timeout=600)
    if not out.exists():
        raise SystemExit(f"!! DWG not written from {dxf}")
    dwg = Path(dxf).with_suffix(".dwg")
    shutil.copy2(out, dwg)
    return dwg


if __name__ == "__main__":
    import sys
    a = sys.argv[1:]
    if a[:1] == ["template"] and len(a) == 2:
        print("wrote", *template(a[1]))
    elif a[:1] == ["draw"] and len(a) == 3:
        m, s = read_xlsx(a[1])
        print("wrote", to_dxf(m, s, a[2]))
    else:
        raise SystemExit(__doc__)
