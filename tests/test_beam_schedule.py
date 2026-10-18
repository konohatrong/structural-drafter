"""drafter.beam_schedule: the spreadsheet notation (user rules 2026-10-18, docs/concrete/BEAM_SCHEDULE_RULES.md), the
round trip and the drawing"""
import pytest

from drafter import beam_schedule as BS

N, X = (lambda n: {"n": n, "x": 0}), (lambda x: {"n": 0, "x": x})


@pytest.mark.parametrize("s, layers, label", [
    ("3/3/3", [N(3), N(3), N(3)], "9 - DB16"),                 # 3 layers, nominal, no extra
    ("3+2", [N(3), X(2)], "3+2 - DB16"),                       # 1 nominal layer + 1 extra layer
    ("3/3+2", [N(3), N(3), X(2)], "6+2 - DB16"),               # 2 nominal layers + 1 extra layer
    ("3+3/2", [N(3), X(3), X(2)], "3+5 - DB16"),               # extra bars in two layers
    ("2(+1)", [{"n": 2, "x": 1}], "2+1 - DB16"),               # extra sharing the nominal layer
    ("3/2(+1)+2", [N(3), {"n": 2, "x": 1}, X(2)], "5+3 - DB16"),
    ("2", [N(2)], "2 - DB16"),
    (" 3 / 3 + 2 ", [N(3), N(3), X(2)], "6+2 - DB16"),
])
def test_bar_grammar(s, layers, label):
    assert BS.parse_bars(s) == layers
    assert BS.bars_label(layers, 16) == label
    assert BS.parse_bars(BS.fmt_bars(layers)) == layers          # round trip through the canonical form


def test_bar_grammar_details():
    assert BS.parse_bars(3) == [N(3)] and BS.parse_bars("3.0") == [N(3)]     # Excel numbers
    assert BS.parse_bars("-") == [] and BS.parse_bars(None) == []
    assert BS.totals(BS.parse_bars("3/2(+1)+2")) == [3, 3, 2]
    assert BS.fmt_bars(BS.parse_bars("3 + 2")) == "3+2"
    for bad in ("3+", "+2", "3/0", "0(+2)", "3+2(+1)", "3x2", "3//2", "3+0"):
        with pytest.raises(ValueError):
            BS.parse_bars(bad)


def test_stirrup_notation():
    assert BS.parse_stirrup("2-RB9@100") == {"legs": 2, "mark": "RB", "db": 9, "s": 100}
    assert BS.parse_stirrup("3-db10 @ 150") == {"legs": 3, "mark": "DB", "db": 10, "s": 150}
    assert BS.parse_stirrup("RB9@200")["legs"] == 2              # legs default: one hoop
    assert BS.fmt_stirrup(BS.parse_stirrup("4-RB9@100")) == "4-RB9@100"
    with pytest.raises(ValueError):
        BS.parse_stirrup("1-RB9@100")
    with pytest.raises(ValueError):
        BS.parse_stirrup("RB9-100")


def test_xlsx_round_trip(tmp_path):
    p = tmp_path / "s.xlsx"
    BS.write_xlsx(p, BS.EXAMPLE)
    marks, settings = BS.read_xlsx(p)
    assert list(marks) == list(BS.EXAMPLE)
    for mk, m in BS.EXAMPLE.items():
        assert (marks[mk]["b"], marks[mk]["h"], marks[mk]["slab"]) == (m["b"], m["h"], m["slab"])
        for tag, c in m["cols"].items():
            r = marks[mk]["cols"][tag]
            assert r["top"] == c["top"] and r["bot"] == c["bot"] and r["stirrup"] == c["stirrup"]
    assert settings["cover"] == 40


def test_xlsx_sheets_and_blank_rows(tmp_path):
    import openpyxl
    p = tmp_path / "s.xlsx"
    BS.write_xlsx(p, {"B1": BS.EXAMPLE["B1"]})
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["BEAMS", "INSTRUCTIONS", "SETTINGS"]
    text = " ".join(str(c.value) for r in wb["INSTRUCTIONS"].iter_rows() for c in r if c.value)
    assert "3/3+2" in text and "2(+1)" in text and "3-RB9@100" in text
    assert wb["BEAMS"]["A4"].value is None                      # second row of B1: mark left blank
    marks, _ = BS.read_xlsx(p)
    assert set(marks["B1"]["cols"]) == {"END SUPPORT", "MID SPAN", "CONTINUOUS SUPPORT"}


def test_bad_cell_names_its_row(tmp_path):
    import openpyxl
    p = tmp_path / "s.xlsx"
    BS.write_xlsx(p, {"B1": BS.EXAMPLE["B1"]})
    wb = openpyxl.load_workbook(p)
    wb["BEAMS"]["F3"] = "3x2"
    wb.save(p)
    with pytest.raises(ValueError, match="row 3"):
        BS.read_xlsx(p)


def test_inner_legs_on_shared_bars():
    xs5 = BS._bar_x(5, 0, 200)
    xs3 = BS._bar_x(3, 0, 200)
    assert BS._inner_legs(xs5, xs3, 1) == [100]                 # crosstie on the middle bar of both layers
    assert BS._inner_legs(BS._bar_x(4, 0, 300), BS._bar_x(4, 0, 300), 2) == [100, 200]
    assert BS._inner_legs(BS._bar_x(3, 0, 200), BS._bar_x(2, 0, 200), 1) is None   # no shared inner bar


def test_draw_example(tmp_path):
    from drafter import td_engine
    n0 = len(td_engine.msp)
    p = tmp_path / "e.xlsx"
    BS.write_xlsx(p, BS.EXAMPLE)
    out = BS.to_dxf(*BS.read_xlsx(p), tmp_path / "t.dxf")
    assert out.exists()
    ents = list(td_engine.msp)[n0:]
    texts = {e.dxf.text for e in ents if e.dxftype() == "TEXT"}
    assert {"B1", "B2", "B1C", "RB1", "3+2 - DB20", "2+1 - DB20", "8+2 - DB20", "3-RB9@100", "2-DB12"} <= texts
    assert not td_engine.WARNINGS
